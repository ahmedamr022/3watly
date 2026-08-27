"""
MAJRA Real CV Parser & ATS Diagnostics Engine
Extracts plain text from PDF, DOCX, and TXT files, categorizes resume sections,
and computes a comprehensive ATS Compatibility Score with actionable recommendations.
"""

import io
import re
from typing import Dict, Any, List, Tuple
from pypdf import PdfReader
import docx

from app.services.nlp_extractor import NLPSkillExtractor


class CVParser:
    """
    Parses CV documents (PDF, DOCX, TXT) and computes ATS diagnostic metrics.
    """

    def __init__(self, nlp_extractor: NLPSkillExtractor = None):
        self.nlp = nlp_extractor or NLPSkillExtractor()

        # Section keywords for detection
        self.section_patterns = {
            "Contact Information": [
                r"\b(email|phone|mobile|tel|linkedin|github|portfolio|address|location)\b",
                r"[\w\.-]+@[\w\.-]+\.\w+",
                r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}"
            ],
            "Work Experience": [
                r"\b(experience|work experience|employment history|professional experience|career history|work history)\b"
            ],
            "Education": [
                r"\b(education|academic background|university|bachelor|master|degree|bsc|msc|phd|diploma)\b"
            ],
            "Technical Skills": [
                r"\b(skills|technical skills|technologies|tools & technologies|core competencies|expertise|proficiencies)\b"
            ],
            "Projects & Portfolio": [
                r"\b(projects|key projects|academic projects|personal projects|portfolio|case studies)\b"
            ],
            "Certifications & Awards": [
                r"\b(certifications|certificates|courses|awards|honors|achievements|training)\b"
            ]
        }

        # Strong action verbs for impact score
        self.action_verbs = [
            "developed", "built", "implemented", "designed", "architected", "engineered",
            "optimized", "improved", "increased", "decreased", "reduced", "led", "managed",
            "deployed", "integrated", "automated", "created", "analyzed", "scaled", "delivered"
        ]

    def extract_text_from_bytes(self, file_bytes: bytes, filename: str, content_type: str = "") -> str:
        """
        Extracts clean plain text from PDF, DOCX, or plain text binary buffers.
        """
        lower_name = filename.lower()
        extracted_text = ""

        # 1. Handle PDF
        if lower_name.endswith(".pdf") or "pdf" in content_type:
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                pages_text = []
                for page in reader.pages:
                    t = page.extract_text()
                    if t:
                        pages_text.append(t)
                extracted_text = "\n".join(pages_text)
            except Exception as e:
                extracted_text = f"[PDF Parsing Fallback Error: {e}]"

        # 2. Handle DOCX
        elif lower_name.endswith(".docx") or "wordprocessingml" in content_type:
            try:
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                extracted_text = "\n".join(paragraphs)
            except Exception as e:
                extracted_text = f"[DOCX Parsing Fallback Error: {e}]"

        # 3. Handle Plain Text (.txt, .md, etc.)
        else:
            try:
                extracted_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                try:
                    extracted_text = file_bytes.decode("latin-1")
                except Exception:
                    extracted_text = str(file_bytes)

        return self._clean_text(extracted_text)

    def _clean_text(self, text: str) -> str:
        """Cleans and standardizes whitespace and line breaks."""
        if not text:
            return ""
        # Normalize multiple newlines and spaces
        cleaned = re.sub(r"\r\n", "\n", text)
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        cleaned = re.sub(r"\n\s*\n+", "\n\n", cleaned)
        return cleaned.strip()

    def analyze_cv(self, raw_text: str) -> Dict[str, Any]:
        """
        Runs NLP skill extraction, detects structural sections, and calculates ATS score.
        """
        words = re.findall(r"\b\w+\b", raw_text.lower())
        word_count = len(words)

        # 1. NLP Skill Extraction
        nlp_res = self.nlp.extract_skills_from_text(raw_text)
        detected_skills = nlp_res.get("all_detected_skills", [])
        canonical_skills = sorted(list({s["name"] for s in detected_skills}))

        # 2. Detect Sections
        detected_sections: List[str] = []
        lower_text = raw_text.lower()

        for section_name, patterns in self.section_patterns.items():
            for p in patterns:
                if re.search(p, lower_text):
                    detected_sections.append(section_name)
                    break

        critical_sections = ["Contact Information", "Work Experience", "Education", "Technical Skills"]
        missing_critical = [s for s in critical_sections if s not in detected_sections]

        # 3. Calculate Component Scores (Max 100)
        # A. Structure & Formatting Score (Max 25 pts)
        structure_score = 0
        if "Contact Information" in detected_sections:
            structure_score += 6
        if "Work Experience" in detected_sections:
            structure_score += 7
        if "Education" in detected_sections:
            structure_score += 6
        if "Technical Skills" in detected_sections:
            structure_score += 6

        # B. Skills Density & Categorization Score (Max 40 pts)
        skills_count = len(canonical_skills)
        if skills_count >= 12:
            skills_score = 40
        elif skills_count >= 8:
            skills_score = 32
        elif skills_count >= 5:
            skills_score = 24
        elif skills_count >= 2:
            skills_score = 15
        else:
            skills_score = max(5, skills_count * 5)

        # C. Impact & Action Verbs Score (Max 20 pts)
        action_verb_count = sum(1 for verb in self.action_verbs if re.search(rf"\b{verb}\b", lower_text))
        has_metrics = bool(re.search(r"\b\d+[%+xX]?\b", raw_text))

        impact_score = min(14, action_verb_count * 2)
        if has_metrics:
            impact_score += 6
        impact_score = min(20, impact_score)

        # D. Readability & Length Score (Max 15 pts)
        if 350 <= word_count <= 900:
            readability_score = 15
        elif 200 <= word_count < 350:
            readability_score = 10
        elif word_count > 900:
            readability_score = 11
        else:
            readability_score = 5

        total_ats_score = min(100, structure_score + skills_score + impact_score + readability_score)

        # 4. Generate Actionable Suggestions
        suggestions = []
        if missing_critical:
            suggestions.append(f"Add missing standard sections: {', '.join(missing_critical)}.")
        if skills_count < 8:
            suggestions.append("Enrich your technical skills section with more canonical tools and frameworks.")
        if not has_metrics:
            suggestions.append("Quantify your achievements with measurable numbers, percentages, or scale metrics.")
        if action_verb_count < 4:
            suggestions.append("Use stronger action verbs (e.g., Developed, Architected, Optimized, Deployed) at the start of bullet points.")
        if word_count < 300:
            suggestions.append("Your CV is too concise; elaborate more on your project responsibilities and accomplishments.")

        if not suggestions:
            suggestions.append("Great job! Your CV has excellent ATS structure, keyword richness, and impact formatting.")

        return {
            "overall_score": total_ats_score,
            "structure_score": structure_score,
            "skills_score": skills_score,
            "impact_score": impact_score,
            "readability_score": readability_score,
            "detected_sections": detected_sections,
            "missing_critical_sections": missing_critical,
            "word_count": word_count,
            "suggestions": suggestions,
            "extracted_skills": canonical_skills,
            "detailed_skills": detected_skills
        }
